import asyncio
from pathlib import Path

import pytest
from pptx import Presentation

from presentation_gen import server


def _sample_slides():
    return [{"title": "Overview", "content": "One\nTwo"}]


def test_create_presentation_writes_new_pptx(tmp_path):
    output = tmp_path / "deck.pptx"

    result = server._create_presentation(
        "Quarterly report", _sample_slides(), output_path=str(output)
    )

    assert result == str(output)
    assert output.is_file()
    assert len(Presentation(output).slides) == 2


def test_create_presentation_refuses_existing_output(tmp_path):
    output = tmp_path / "deck.pptx"
    output.write_bytes(b"keep-me")

    with pytest.raises(FileExistsError, match="already exists"):
        server._create_presentation(
            "Quarterly report", _sample_slides(), output_path=str(output)
        )

    assert output.read_bytes() == b"keep-me"


def test_create_presentation_does_not_overwrite_concurrently_created_output(
    tmp_path, monkeypatch
):
    output = tmp_path / "deck.pptx"
    original_new_output_path = server._new_output_path

    def create_output_after_validation(*args, **kwargs):
        filepath = original_new_output_path(*args, **kwargs)
        filepath.write_bytes(b"keep-me")
        return filepath

    monkeypatch.setattr(server, "_new_output_path", create_output_after_validation)

    with pytest.raises(FileExistsError):
        server._create_presentation(
            "Quarterly report", _sample_slides(), output_path=str(output)
        )

    assert output.read_bytes() == b"keep-me"


def test_new_file_writer_removes_partial_output_on_failure(tmp_path):
    output = tmp_path / "partial.pptx"

    def fail_after_write(output_file):
        output_file.write(b"partial")
        raise RuntimeError("serialization failed")

    with pytest.raises(RuntimeError, match="serialization failed"):
        server._write_new_file(output, fail_after_write)

    assert not output.exists()


def test_create_presentation_requires_pptx_extension(tmp_path):
    with pytest.raises(ValueError, match=r"\.pptx"):
        server._create_presentation(
            "Quarterly report",
            _sample_slides(),
            output_path=str(tmp_path / "deck.pdf"),
        )


def test_create_presentation_rejects_excessive_slide_count():
    slides = [
        {"title": f"Slide {index}"} for index in range(server.MAX_CONTENT_SLIDES + 1)
    ]

    with pytest.raises(ValueError, match="at most"):
        server._create_presentation("Large deck", slides)


def test_thumbnail_rejects_invalid_width_before_allocating(tmp_path):
    source = tmp_path / "deck.pptx"
    Presentation().save(source)

    with pytest.raises(ValueError, match="width"):
        server._create_thumbnail(str(source), server.MAX_THUMBNAIL_WIDTH + 1)


def test_thumbnail_refuses_existing_output(tmp_path):
    source = tmp_path / "deck.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)
    output = tmp_path / "preview.png"
    output.write_bytes(b"keep-me")

    with pytest.raises(FileExistsError, match="already exists"):
        server._create_thumbnail(str(source), output_path=str(output))

    assert output.read_bytes() == b"keep-me"


def test_export_refuses_existing_output_before_invoking_libreoffice(tmp_path):
    source = tmp_path / "deck.pptx"
    Presentation().save(source)
    output = tmp_path / "deck.pdf"
    output.write_bytes(b"keep-me")

    with pytest.raises(FileExistsError, match="already exists"):
        server._export_to_pdf(str(source), str(output))

    assert output.read_bytes() == b"keep-me"


def test_export_converts_in_temporary_directory(tmp_path, monkeypatch):
    source = tmp_path / "deck.pptx"
    Presentation().save(source)
    output = tmp_path / "final.pdf"

    monkeypatch.setattr(server.shutil, "which", lambda _: "/usr/bin/soffice")

    def fake_run(command, **kwargs):
        out_dir = Path(command[command.index("--outdir") + 1])
        assert out_dir != tmp_path
        (out_dir / "deck.pdf").write_bytes(b"pdf-data")
        return type("Result", (), {"returncode": 0, "stderr": "", "stdout": ""})()

    monkeypatch.setattr(server.subprocess, "run", fake_run)

    assert server._export_to_pdf(str(source), str(output)) == str(output)
    assert output.read_bytes() == b"pdf-data"


def test_thumbnail_writes_new_png(tmp_path):
    source = tmp_path / "deck.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)
    output = tmp_path / "preview.png"

    assert server._create_thumbnail(str(source), 640, str(output)) == str(output)
    assert output.is_file()


def test_tool_schemas_publish_resource_limits():
    tools = {tool.name: tool for tool in asyncio.run(server.handle_list_tools())}

    slides = tools["create_slides"].inputSchema["properties"]["slides"]
    width = tools["create_thumbnail"].inputSchema["properties"]["width"]

    assert slides["maxItems"] == server.MAX_CONTENT_SLIDES
    assert width["minimum"] == server.MIN_THUMBNAIL_WIDTH
    assert width["maximum"] == server.MAX_THUMBNAIL_WIDTH
